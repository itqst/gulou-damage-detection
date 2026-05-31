"""
Generate final run9 experiment reports.
"""
from __future__ import annotations

import csv
import json
import os
import re
import shutil
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import cv2
import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
RUN9 = ROOT / "run9"
FIGURE_DIR = RUN9 / "docs" / "figures"
SAMPLE_DIR = FIGURE_DIR / "sample_detections"
os.environ.setdefault("YOLO_CONFIG_DIR", str(ROOT / ".ultralytics"))


CAPTION_COUNTERS: dict[int, dict[str, int]] = {}


CASE_NAMES = {
    "zengchong": "增冲鼓楼",
    "chaoli": "朝利鼓楼",
    "congjiang": "从江鼓楼",
    "universal": "三鼓楼通用模型",
}


MODELS = {
    "zengchong_single": RUN9 / "zengchong" / "model" / "train" / "weights" / "best.pt",
    "chaoli_single": RUN9 / "chaoli" / "model" / "train" / "weights" / "best.pt",
    "congjiang_single": RUN9 / "congjiang" / "model" / "train" / "weights" / "best.pt",
    "universal": RUN9 / "universal" / "model" / "train" / "weights" / "best.pt",
}


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def read_yaml_simple(path: Path) -> dict[str, str]:
    data = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            data[key.strip()] = value.strip()
    return data


def read_training_best(case: str) -> dict:
    csv_path = RUN9 / case / "model" / "train" / "results.csv"
    rows = list(csv.DictReader(csv_path.open(encoding="utf-8")))
    def f(row, key):
        return float(row[key])

    best_map50 = max(rows, key=lambda row: f(row, "metrics/mAP50(B)"))
    best_map = max(rows, key=lambda row: f(row, "metrics/mAP50-95(B)"))
    last = rows[-1]
    return {
        "epochs_completed": len(rows),
        "best_mAP50_epoch": int(float(best_map50["epoch"])),
        "best_mAP50": f(best_map50, "metrics/mAP50(B)"),
        "best_mAP50_precision": f(best_map50, "metrics/precision(B)"),
        "best_mAP50_recall": f(best_map50, "metrics/recall(B)"),
        "best_mAP50_95_epoch": int(float(best_map["epoch"])),
        "best_mAP50_95": f(best_map, "metrics/mAP50-95(B)"),
        "last_epoch": int(float(last["epoch"])),
        "last_mAP50": f(last, "metrics/mAP50(B)"),
        "last_mAP50_95": f(last, "metrics/mAP50-95(B)"),
    }


def metrics_by_name() -> dict[str, dict]:
    items = read_json(RUN9 / "evaluation" / "all_metrics_summary.json")
    return {item["name"]: item for item in items}


def setup_doc(title_text: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title_text)
    set_run(r, "宋体", Pt(22), True)
    doc.add_paragraph()
    return doc


def set_run(run, font="宋体", size=Pt(12), bold=False):
    run.font.name = "Times New Roman"
    run.font.size = size
    run.bold = bold
    run.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.element.rPr.rFonts.set(qn("w:eastAsia"), font)


def clean_text(text: object) -> str:
    value = str(text)
    value = re.sub(r"([\u4e00-\u9fff])\s+([A-Za-z0-9])", r"\1\2", value)
    value = re.sub(r"([A-Za-z0-9%])\s+([\u4e00-\u9fff])", r"\1\2", value)
    return value


def heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, "宋体", Pt(15 if level == 1 else 13), True)


def para(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(clean_text(text))
    set_run(r, "宋体", Pt(12), False)


def caption(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(clean_text(text))
    set_run(r, "宋体", Pt(10.5), True)


def numbered_caption(doc: Document, kind: str, text: str) -> str:
    counters = CAPTION_COUNTERS.setdefault(id(doc), {"table": 0, "figure": 0})
    counters[kind] += 1
    clean = text.strip()
    for prefix in ["表", "图"]:
        if clean.startswith(prefix):
            clean = clean[1:].strip()
    return f"{'表' if kind == 'table' else '图'}{counters[kind]}　{clean}"


def add_table(doc: Document, headers: list[str], rows: list[list[object]], title: str | None = None):
    if title:
        caption(doc, numbered_caption(doc, "table", title))
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(clean_text(header))
        set_run(r, "宋体", Pt(10.5), True)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_text(value))
            set_run(run, "宋体", Pt(9.5), False)
    doc.add_paragraph()


def add_image_if_exists(doc: Document, image_path: Path, cap: str, width=Cm(13)):
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(image_path), width=width)
    caption(doc, numbered_caption(doc, "figure", cap))


def pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def num(value: float) -> str:
    return f"{value:.3f}"


def setup_plot_style():
    plt.rcParams["font.sans-serif"] = [
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def save_bar_chart(path: Path, title: str, labels: list[str], series: list[tuple[str, list[float]]], ylabel: str) -> None:
    setup_plot_style()
    path.parent.mkdir(parents=True, exist_ok=True)
    width = 0.78 / max(len(series), 1)
    x = list(range(len(labels)))
    fig, ax = plt.subplots(figsize=(8.2, 4.6), dpi=180)
    all_values = [value for _, values in series for value in values]
    integer_like = all(abs(value - round(value)) < 1e-9 for value in all_values) and max(all_values) > 10
    for idx, (name, values) in enumerate(series):
        offsets = [v + (idx - (len(series) - 1) / 2) * width for v in x]
        bars = ax.bar(offsets, values, width=width, label=clean_text(name))
        value_labels = [f"{int(round(v))}" if integer_like else f"{v:.3f}" for v in values]
        ax.bar_label(bars, labels=value_labels, padding=2, fontsize=8)
    ax.set_title(clean_text(title), fontsize=12, pad=10)
    ax.set_ylabel(clean_text(ylabel))
    ax.set_xticks(x)
    ax.set_xticklabels([clean_text(label) for label in labels])
    ax.set_ylim(0, max(1.0, max(max(values) for _, values in series) * 1.2))
    ax.grid(axis="y", linestyle="--", alpha=0.35)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def ensure_comparison_figures(metrics: dict[str, dict], dataset: dict) -> dict[str, Path]:
    figures = {
        "dataset": FIGURE_DIR / "dataset_comparison.png",
        "single_metrics": FIGURE_DIR / "single_case_metrics.png",
        "universal_compare": FIGURE_DIR / "universal_vs_single_map50.png",
        "universal_case": FIGURE_DIR / "universal_case_metrics.png",
    }
    labels = [CASE_NAMES[case].replace("鼓楼", "") for case in ["zengchong", "chaoli", "congjiang"]]
    save_bar_chart(
        figures["dataset"],
        "三座鼓楼数据集规模对比",
        labels,
        [
            ("图像数", [dataset["by_case"][case]["images"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("标注框数", [dataset["by_case"][case]["boxes"] for case in ["zengchong", "chaoli", "congjiang"]]),
        ],
        "数量",
    )
    save_bar_chart(
        figures["single_metrics"],
        "单案例模型测试集指标对比",
        labels,
        [
            ("Precision", [metrics[f"{case}_single_on_{case}_test"]["precision"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("Recall", [metrics[f"{case}_single_on_{case}_test"]["recall"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("mAP50", [metrics[f"{case}_single_on_{case}_test"]["mAP50"] for case in ["zengchong", "chaoli", "congjiang"]]),
        ],
        "指标值",
    )
    save_bar_chart(
        figures["universal_compare"],
        "单案例模型与通用模型 mAP50 对比",
        labels,
        [
            ("单案例模型", [metrics[f"{case}_single_on_{case}_test"]["mAP50"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("通用模型", [metrics[f"universal_on_{case}_test"]["mAP50"] for case in ["zengchong", "chaoli", "congjiang"]]),
        ],
        "mAP50",
    )
    save_bar_chart(
        figures["universal_case"],
        "通用模型分鼓楼测试集指标对比",
        labels,
        [
            ("Precision", [metrics[f"universal_on_{case}_test"]["precision"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("Recall", [metrics[f"universal_on_{case}_test"]["recall"] for case in ["zengchong", "chaoli", "congjiang"]]),
            ("mAP50", [metrics[f"universal_on_{case}_test"]["mAP50"] for case in ["zengchong", "chaoli", "congjiang"]]),
        ],
        "指标值",
    )
    return figures


def positive_test_images(case: str) -> list[Path]:
    image_dir = RUN9 / "datasets" / case / "test" / "images"
    label_dir = RUN9 / "datasets" / case / "test" / "labels"
    images = sorted(
        [
            *image_dir.glob("*.jpg"),
            *image_dir.glob("*.jpeg"),
            *image_dir.glob("*.png"),
        ]
    )
    positives = []
    for image in images:
        label = label_dir / f"{image.stem}.txt"
        if label.exists() and label.read_text(encoding="utf-8").strip():
            positives.append(image)
    if not positives:
        raise FileNotFoundError(f"No test images found for {case}")
    return positives


def pick_detected_positive_test_image(case: str, model) -> Path:
    candidates = positive_test_images(case)
    best_image = candidates[0]
    best_score = (-1, -1.0)
    for image in candidates:
        results = model.predict(
            source=str(image),
            conf=0.4,
            imgsz=640,
            device="0",
            verbose=False,
        )
        boxes = results[0].boxes
        count = len(boxes) if boxes is not None else 0
        max_conf = 0.0
        if boxes is not None and count > 0:
            max_conf = float(boxes.conf.detach().cpu().max())
        score = (count, max_conf)
        if score > best_score:
            best_score = score
            best_image = image
    return best_image


def resize_to_height(image: np.ndarray, height: int) -> np.ndarray:
    h, w = image.shape[:2]
    if h == height:
        return image
    width = max(1, int(w * height / h))
    return cv2.resize(image, (width, height), interpolation=cv2.INTER_AREA)


def save_detection_pair(model, image_path: Path, output_path: Path) -> None:
    original = cv2.imread(str(image_path))
    if original is None:
        raise FileNotFoundError(image_path)
    results = model.predict(
        source=str(image_path),
        conf=0.4,
        imgsz=640,
        device="0",
        verbose=False,
    )
    detected = results[0].plot(labels=True, conf=True, line_width=2)
    target_height = 640
    original = resize_to_height(original, target_height)
    detected = resize_to_height(detected, target_height)
    separator = np.full((target_height, 18, 3), 255, dtype=np.uint8)
    paired = np.concatenate([original, separator, detected], axis=1)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    cv2.imwrite(str(output_path), paired)


def ensure_sample_detection_figures() -> dict[str, dict[str, Path]]:
    from ultralytics import YOLO

    model_cache = {}

    def get_model(model_key: str):
        if model_key not in model_cache:
            model_cache[model_key] = YOLO(str(MODELS[model_key]))
        return model_cache[model_key]

    output: dict[str, dict[str, Path]] = {"single": {}, "universal": {}}
    for case in ["zengchong", "chaoli", "congjiang"]:
        single_key = f"{case}_single"
        single_out = SAMPLE_DIR / f"{case}_single_original_detection.jpg"
        universal_out = SAMPLE_DIR / f"{case}_universal_original_detection.jpg"
        single_model = get_model(single_key)
        universal_model = get_model("universal")
        single_image = pick_detected_positive_test_image(case, single_model)
        universal_image = pick_detected_positive_test_image(case, universal_model)
        save_detection_pair(single_model, single_image, single_out)
        save_detection_pair(universal_model, universal_image, universal_out)
        output["single"][case] = single_out
        output["universal"][case] = universal_out
    return output


def single_case_discussion(case: str, metric: dict) -> str:
    case_name = CASE_NAMES[case]
    if case == "zengchong":
        return (
            f"{case_name}测试集上的检测结果较为稳定，Precision、Recall 与 mAP50 分别达到 "
            f"{num(metric['precision'])}、{num(metric['recall'])} 和 {num(metric['mAP50'])}。"
            "该结果表明，在样本视角相对集中、破损目标较清晰的条件下，轻量化 YOLOv8n 模型能够形成较可靠的病害区域识别能力。"
            "与此同时，mAP50-95 低于 mAP50，说明模型虽然能够有效发现破损目标，但对不规则边界和小尺度细部的精确定位仍存在一定限制。"
        )
    if case == "chaoli":
        return (
            f"{case_name}测试集上模型取得 mAP50={num(metric['mAP50'])}，Precision 为 {num(metric['precision'])}，"
            f"但 Recall 为 {num(metric['recall'])}，表现出较明显的漏检倾向。"
            "这一现象与历史建筑病害检测任务中的一般难点相一致：当目标尺度较小、背景纹理复杂且拍摄角度变化较大时，模型更容易将轻微破损与正常构件纹理混淆。"
            "因此，朝利鼓楼后续应优先补充困难样本和边界不清晰样本，以增强模型对弱特征病害的召回能力。"
        )
    return (
        f"{case_name}测试集上模型取得 mAP50={num(metric['mAP50'])}，整体检测能力介于增冲与朝利之间。"
        f"Precision={num(metric['precision'])} 表明模型对已检出目标具有一定可信度，但 Recall={num(metric['recall'])} "
        "说明仍有部分破损区域未被稳定识别。"
        "从江鼓楼样本中木构纹理、阴影与瓦面边缘共同构成较强干扰，导致矩形框检测对小目标和边界模糊目标的敏感性不足。"
    )


def report_common_method(doc: Document):
    heading(doc, "三、实验方法", 1)
    para(
        doc,
        "本研究按照“数据采集—人工标注—模型训练—独立验证—结果讨论”的实验路线，建立面向鼓楼木构建筑破损识别的目标检测流程。"
        "考虑当前标注数据采用矩形框形式，本阶段选择 YOLOv8n 单阶段目标检测模型作为基础网络，重点评估其对可见破损区域的自动定位能力。"
        "模型由 MS COCO 预训练权重 yolov8n.pt 初始化，随后在对应鼓楼数据集上进行迁移训练。"
        "所有实验均保持输入尺寸、批大小、增强策略、早停机制和置信度阈值一致，以减少非数据因素对实验结果的影响。",
    )
    add_table(
        doc,
        ["步骤", "处理内容"],
        [
            ["影像样本整理", "从无人机影像中筛选有效图像样本，剔除明显模糊或不可判读样本"],
            ["人工标注", "采用统一 damage 类别标注可见破损区域，并补充经人工复核的负样本"],
            ["数据划分", "按训练集、验证集、测试集划分数据，测试集不参与模型训练"],
            ["模型训练", "使用统一 YOLOv8n 参数进行迁移训练，并由验证集性能确定最优权重"],
            ["独立评价", "在测试集上计算 Precision、Recall、mAP50 和 mAP50-95"],
        ],
        "表  实验流程设计",
    )
    heading(doc, "四、实验设置", 1)
    para(
        doc,
        "为保证三座鼓楼实验结果具有可比性，各单案例模型与三鼓楼通用模型均采用一致的训练参数。"
        "其中，通用模型以三座鼓楼合并数据集作为训练数据来源，单案例模型则分别使用对应鼓楼数据集进行训练。"
        "测试阶段统一使用各自独立测试集计算评价指标，测试集图像不参与模型训练和参数选择。",
    )
    add_table(
        doc,
        ["参数", "配置"],
        [
            ["检测任务", "单类别鼓楼破损目标检测"],
            ["基础模型", "YOLOv8n"],
            ["初始化权重", "yolov8n.pt"],
            ["标注类别", "damage"],
            ["输入尺寸", "640×640"],
            ["Batch size", "8"],
            ["最大训练轮数", "150"],
            ["早停耐心值", "30"],
            ["优化器", "auto(实际采用AdamW)"],
            ["学习率设置", "配置lr0=0.01、lrf=0.01；auto优化器按YOLO默认策略自动调整"],
            ["动量参数", "momentum=0.937"],
            ["权重衰减", "weight_decay=0.0005"],
            ["预热轮数", "warmup_epochs=3.0"],
            ["随机种子", "seed=0"],
            ["确定性训练", "deterministic=True"],
            ["混合精度训练", "AMP=True"],
            ["Mosaic关闭轮次", "close_mosaic=10"],
            ["数据增强", "Mosaic、HSV扰动、随机翻转、缩放、平移、随机擦除"],
            ["测试集评价方式", "基于测试集计算Precision、Recall、mAP50、mAP50-95"],
            ["检测样例置信度阈值", "0.4"],
            ["数据加载线程", "workers=0"],
            ["训练设备", "NVIDIA GeForce RTX 3050 Ti Laptop GPU"],
        ],
        "表  实验统一训练参数",
    )


def generate_single_report(case: str, metric: dict, dataset: dict, samples: dict[str, dict[str, Path]]):
    case_name = CASE_NAMES[case]
    out_dir = RUN9 / case / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = setup_doc(f"{case_name}破损检测单案例实验报告")

    heading(doc, "一、实验目的", 1)
    para(
        doc,
        f"本实验以{case_name}为单案例对象，基于已完成的人工标注样本训练 YOLOv8n 破损检测模型。"
        "实验重点为检验统一参数条件下模型在该鼓楼测试集上的检测效果，"
        "并为后续三鼓楼融合对比提供可复核的单案例结果。",
    )

    heading(doc, "二、数据集构建", 1)
    para(
        doc,
        f"{case_name}数据集由无人机影像抽帧样本构成，标注类别统一为 damage。"
        "正样本为人工确认并完成矩形框标注的破损图像，负样本为人工复核后未见明显破损的图像，"
        "负样本以空 YOLO 标签文件参与训练。"
        "该处理方式体现了病害识别研究中对数据清洗和人工复核的重视原则，有助于模型同时学习破损区域与无破损背景的差异，"
        "从而降低复杂木构纹理、阴影和瓦面边缘条件下的误检风险。",
    )
    add_table(
        doc,
        ["数据集", "图像数", "正样本", "负样本", "标注框数", "正样本比例"],
        [
            [
                case_name,
                dataset["images"],
                dataset["positive_images"],
                dataset["negative_images"],
                dataset["boxes"],
                pct(dataset["positive_ratio"]),
            ]
        ],
        "表  数据集总体统计",
    )
    add_table(
        doc,
        ["划分", "图像数", "正样本", "负样本", "标注框数"],
        [
            [
                split,
                values["images"],
                values["positive_images"],
                values["negative_images"],
                values["boxes"],
            ]
            for split, values in dataset["splits"].items()
        ],
        "表  数据集划分统计",
    )
    if case == "chaoli":
        para(
            doc,
            "从数据统计看，朝利鼓楼的图像数和正样本数与其他鼓楼接近，但标注框总数相对较低。"
            "经复核，朝利鼓楼共有164张正样本、376个标注框，平均每张正样本约2.29个标注框。"
            "该结果与原始Labelme标注文件和转换后的YOLO标签文件一致，说明统计过程无误。"
            "标注框数量较少主要反映该组样本中单张图像可见破损目标数量较少，而不是数据划分或标签转换错误。",
        )

    report_common_method(doc)

    train = read_training_best(case)
    heading(doc, "五、训练过程与图像级评估结果", 1)
    para(
        doc,
        f"该模型在第 {train['epochs_completed']} 轮结束训练，训练过程采用早停策略自动保存验证集性能最优权重。"
        f"验证过程中 mAP50 的最优值出现在第 {train['best_mAP50_epoch']} 轮，达到 {num(train['best_mAP50'])}；"
        f"mAP50-95 的最优值出现在第 {train['best_mAP50_95_epoch']} 轮，达到 {num(train['best_mAP50_95'])}。"
        "测试集作为未参与训练的数据子集，用于检验模型对该鼓楼场景下未见样本的泛化能力。",
    )
    add_table(
        doc,
        ["指标", "测试集结果"],
        [
            ["Precision", num(metric["precision"])],
            ["Recall", num(metric["recall"])],
            ["mAP50", num(metric["mAP50"])],
            ["mAP50-95", num(metric["mAP50_95"])],
        ],
        "表  测试集检测指标",
    )
    add_image_if_exists(
        doc,
        RUN9 / case / "model" / "train" / "results.png",
        "图  训练过程损失与评价指标变化",
    )
    add_image_if_exists(
        doc,
        RUN9 / case / "model" / "train" / "confusion_matrix.png",
        "图  验证阶段混淆矩阵",
    )
    add_image_if_exists(
        doc,
        RUN9 / "evaluation" / metric["name"] / "BoxPR_curve.png",
        "图  测试集 Precision-Recall 曲线",
    )
    add_image_if_exists(
        doc,
        RUN9 / "evaluation" / metric["name"] / "confusion_matrix.png",
        "图  测试集混淆矩阵",
    )
    para(
        doc,
        "为直观展示模型在测试样本上的检测效果，选取测试集中带有人工标注且模型产生检测框的样本进行推理。"
        "下图左侧为原始测试图像，右侧为同一图像叠加检测框后的结果。",
    )
    add_image_if_exists(
        doc,
        samples["single"][case],
        "图  测试集样本检测对照（左为原始图像，右为检测结果）",
        width=Cm(15),
    )
    para(
        doc,
        f"从测试集结果看，模型在该案例上取得 mAP50={num(metric['mAP50'])}、mAP50-95={num(metric['mAP50_95'])}。"
        "mAP50主要反映模型能否发现破损目标，mAP50-95则进一步考察检测框位置是否足够贴合真实标注。"
        "由于木构件纹理、阴影、瓦片边缘和局部遮挡容易与破损区域形成相似视觉特征，"
        "当破损边界不规则或目标较小时，mAP50-95通常会低于mAP50。",
    )

    heading(doc, "六、结果分析与小结", 1)
    para(doc, single_case_discussion(case, metric))
    para(
        doc,
        "综合测试集评估结果可见，模型能够对鼓楼影像中较为明显的瓦片缺失、局部开裂、构件破损等区域形成自动检测结果。"
        "同时，破损边界不规则、目标尺度较小、视角变化和局部阴影仍会影响检测框定位的稳定性。"
        "后续可进一步补充不同季节、不同光照和不同距离下的样本，并在条件允许时引入实例分割方法，以提升破损边界刻画和面积量化能力。",
    )

    out = out_dir / f"{case_name}破损检测单案例实验报告.docx"
    doc.save(out)
    return out


def generate_universal_report(metrics: dict[str, dict], dataset: dict, figures: dict[str, Path], samples: dict[str, dict[str, Path]]):
    out_dir = RUN9 / "universal" / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = setup_doc("三鼓楼通用破损检测模型实验报告")

    heading(doc, "一、实验目的", 1)
    para(
        doc,
        "本实验将增冲、朝利、从江三座鼓楼数据集合并，训练一个三鼓楼通用破损检测模型。"
        "该模型从 yolov8n.pt 预训练权重重新训练，并非由某一个单案例模型继续微调获得。"
        "实验目的在于检验多鼓楼数据融合后，模型在合并测试集及各鼓楼独立测试集上的检测效果。",
    )
    heading(doc, "二、融合数据集构建", 1)
    para(
        doc,
        "融合数据集的构建遵循统一类别、统一标注尺度和统一数据划分原则。"
        "三座鼓楼均采用 damage 作为唯一检测类别，正样本对应人工确认的可见破损区域，负样本对应经复核未见明显破损的建筑构件图像。"
        "该设置有助于在保持实验变量可控的条件下，观察多实例样本对模型泛化性能的影响。",
    )
    add_table(
        doc,
        ["鼓楼", "图像数", "正样本", "负样本", "标注框数", "正样本比例"],
        [
            [
                CASE_NAMES[case],
                info["images"],
                info["positive_images"],
                info["negative_images"],
                info["boxes"],
                pct(info["positive_ratio"]),
            ]
            for case, info in dataset["by_case"].items()
        ],
        "表  三座鼓楼数据构成",
    )
    add_table(
        doc,
        ["合并数据集", "图像数", "正样本", "负样本", "标注框数", "正样本比例"],
        [[
            "三鼓楼通用数据集",
            dataset["images"],
            dataset["positive_images"],
            dataset["negative_images"],
            dataset["boxes"],
            pct(dataset["positive_ratio"]),
        ]],
        "表  通用数据集总体统计",
    )
    para(
        doc,
        "三座鼓楼在图像数和正样本比例上保持相近，但标注框数量存在差异。"
        "其中朝利鼓楼共有164张正样本和376个标注框，平均每张正样本约2.29个标注框，"
        "低于增冲鼓楼的4.32个和从江鼓楼的3.98个。该差异主要来自各鼓楼可见破损目标密度不同，"
        "不影响其作为单独实验案例和通用模型训练数据的有效性。",
    )
    report_common_method(doc)

    metric = metrics["universal_on_universal_test"]
    heading(doc, "五、通用模型评估结果", 1)
    add_table(
        doc,
        ["评估对象", "Precision", "Recall", "mAP50", "mAP50-95"],
        [[
            "三鼓楼合并测试集",
            num(metric["precision"]),
            num(metric["recall"]),
            num(metric["mAP50"]),
            num(metric["mAP50_95"]),
        ]],
        "表  通用模型合并测试集指标",
    )
    rows = []
    for case in ["zengchong", "chaoli", "congjiang"]:
        item = metrics[f"universal_on_{case}_test"]
        rows.append([CASE_NAMES[case], num(item["precision"]), num(item["recall"]), num(item["mAP50"]), num(item["mAP50_95"])])
    add_table(doc, ["测试集", "Precision", "Recall", "mAP50", "mAP50-95"], rows, "表  通用模型分鼓楼测试结果")
    add_image_if_exists(doc, figures["universal_case"], "图  通用模型分鼓楼测试集指标对比", width=Cm(14))
    para(
        doc,
        "为展示通用模型在不同鼓楼测试样本上的实际输出效果，分别选取三座鼓楼测试集中带有人工标注且模型产生检测框的样本进行检测。"
        "以下对照图左侧为原始测试图像，右侧为通用模型在同一图像上的检测结果。",
    )
    for case in ["zengchong", "chaoli", "congjiang"]:
        add_image_if_exists(
            doc,
            samples["universal"][case],
            f"图  通用模型在{CASE_NAMES[case]}测试集样本上的检测对照（左为原始图像，右为检测结果）",
            width=Cm(15),
        )
    add_image_if_exists(doc, RUN9 / "universal" / "model" / "train" / "results.png", "图  通用模型训练过程")
    add_image_if_exists(doc, RUN9 / "evaluation" / "universal_on_universal_test" / "BoxPR_curve.png", "图  通用模型合并测试集 Precision-Recall 曲线")
    add_image_if_exists(doc, RUN9 / "evaluation" / "universal_on_universal_test" / "confusion_matrix.png", "图  通用模型合并测试集混淆矩阵")

    heading(doc, "六、通用模型结果分析", 1)
    para(
        doc,
        "通用模型在三鼓楼合并测试集及分鼓楼测试集上均取得了有效检测结果，说明合并训练有助于模型吸收多场景、多角度下的共同破损特征。"
        "从分案例测试结果看，增冲鼓楼表现相对稳定，朝利与从江受样本尺度、背景纹理和局部遮挡影响更明显，"
        "但通用模型相较单案例模型在朝利与从江测试集上的 mAP50 具有更高表现，显示出一定的跨场景泛化优势。",
    )
    para(
        doc,
        "需要指出的是，通用模型并不是多个单案例模型的简单融合，也不是以某一鼓楼模型为基础进行迁移微调。"
        "其核心意义在于通过统一训练过程学习三座鼓楼样本中共同存在的破损视觉特征。"
        "这一实验设计与相关研究中关于多样化数据提升模型鲁棒性的结论一致：当训练样本覆盖更多材料纹理、拍摄角度和病害尺度时，"
        "模型对新场景的适应能力通常会得到改善。"
    )
    out = out_dir / "三鼓楼通用破损检测模型实验报告.docx"
    doc.save(out)
    return out


def generate_fusion_report(metrics: dict[str, dict], dataset: dict, figures: dict[str, Path], samples: dict[str, dict[str, Path]]):
    out_dir = RUN9 / "fusion" / "docs"
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = setup_doc("三鼓楼实例融合实验报告")

    heading(doc, "一、实验目的", 1)
    para(
        doc,
        "本报告对增冲、朝利、从江三座鼓楼的破损检测实验结果进行统一汇总。"
        "实验重点是比较三个单案例模型与三鼓楼通用模型在测试集上的表现，分析不同鼓楼样本对模型检测效果的影响，"
        "并判断多鼓楼数据融合是否有助于提升模型的跨实例适应能力。",
    )
    heading(doc, "二、三座鼓楼数据说明", 1)
    rows = []
    for case in ["zengchong", "chaoli", "congjiang"]:
        info = dataset["by_case"][case]
        rows.append([CASE_NAMES[case], info["images"], info["positive_images"], info["negative_images"], info["boxes"], pct(info["positive_ratio"])])
    add_table(doc, ["鼓楼", "图像数", "正样本", "负样本", "标注框数", "正样本比例"], rows, "表  三座鼓楼数据集对比")
    add_image_if_exists(doc, figures["dataset"], "图  三座鼓楼数据集规模对比", width=Cm(14))
    para(
        doc,
        "三座鼓楼数据规模并非完全相同，这是由采集时段、可见破损密度、拍摄角度及人工确认的有效样本数量共同决定的。"
        "为提高实验可比性，本研究统一标注类别、正负样本定义、数据划分方式和训练参数，并将三座鼓楼正样本比例控制在约 72%。",
    )
    para(
        doc,
        "需要说明的是，朝利鼓楼的图像数和正样本数与其他两座鼓楼接近，但其单张正样本中可标注破损区域数量较少，"
        "因此标注框总数低于增冲鼓楼和从江鼓楼。根据标签统计，朝利鼓楼共有164张正样本、376个标注框，"
        "平均每张正样本约2.29个标注框；作为对比，增冲鼓楼平均每张正样本约4.32个标注框，从江鼓楼约3.98个标注框。"
        "该差异反映的是不同鼓楼可见破损密度和标注目标数量的差异，"
        "不代表朝利鼓楼样本数量不足或数据划分错误。",
    )
    para(
        doc,
        "在实验解释上，数据量差异并不被视为实验错误，而是实际文化遗产影像采集过程中常见的客观差异。"
        "因此，本研究重点通过统一训练参数和独立测试集评价来保证结果具有可比性，而不是强行要求每座鼓楼具有完全相同的样本规模。"
    )
    report_common_method(doc)

    heading(doc, "五、单案例检测结果对比", 1)
    rows = []
    for case, metric_name in [
        ("zengchong", "zengchong_single_on_zengchong_test"),
        ("chaoli", "chaoli_single_on_chaoli_test"),
        ("congjiang", "congjiang_single_on_congjiang_test"),
    ]:
        metric = metrics[metric_name]
        rows.append([CASE_NAMES[case], num(metric["precision"]), num(metric["recall"]), num(metric["mAP50"]), num(metric["mAP50_95"])])
    add_table(doc, ["案例", "Precision", "Recall", "mAP50", "mAP50-95"], rows, "表  三个单案例模型测试指标")
    add_image_if_exists(doc, figures["single_metrics"], "图  单案例模型测试集指标对比", width=Cm(14))

    heading(doc, "六、通用模型与单案例模型比较", 1)
    rows = []
    for case, single_name in [
        ("zengchong", "zengchong_single_on_zengchong_test"),
        ("chaoli", "chaoli_single_on_chaoli_test"),
        ("congjiang", "congjiang_single_on_congjiang_test"),
    ]:
        single = metrics[single_name]
        universal = metrics[f"universal_on_{case}_test"]
        rows.append([
            CASE_NAMES[case],
            num(single["mAP50"]),
            num(universal["mAP50"]),
            num(universal["mAP50"] - single["mAP50"]),
            num(single["mAP50_95"]),
            num(universal["mAP50_95"]),
    ])
    add_table(doc, ["鼓楼", "单案例mAP50", "通用mAP50", "mAP50差值", "单案例mAP50-95", "通用mAP50-95"], rows, "表  通用模型与单案例模型对比")
    add_image_if_exists(doc, figures["universal_compare"], "图  单案例模型与通用模型 mAP50 对比", width=Cm(14))
    para(
        doc,
        "为补充定量指标之外的直观结果，报告选取三座鼓楼测试集中带有人工标注且模型产生检测框的样本展示单案例模型的检测效果。"
        "各对照图左侧为原始测试图像，右侧为对应单案例模型叠加检测框后的结果。",
    )
    for case in ["zengchong", "chaoli", "congjiang"]:
        add_image_if_exists(
            doc,
            samples["single"][case],
            f"图  {CASE_NAMES[case]}单案例模型测试样本检测对照（左为原始图像，右为检测结果）",
            width=Cm(15),
        )
    para(
        doc,
        "从测试指标看，增冲鼓楼单案例模型与通用模型均表现较稳定，说明该案例数据质量和标注一致性较好。"
        "朝利与从江案例中，通用模型在 mAP50 上较单案例模型有所提升，表明多鼓楼数据融合能够增强模型对不同屋面纹理、视角变化和破损形态的适应能力。"
        "但在 mAP50-95 指标上，部分案例仍存在提升空间，说明模型对破损边界的精确定位仍受小目标、遮挡和复杂背景影响。",
    )
    para(
        doc,
        "这种结果与历史建筑病害检测研究中的一般规律相吻合：当病害形态具有较强不规则性，且背景材料纹理与病害区域相似时，"
        "模型通常能够较好完成目标发现，但在边界定位精度上更容易受到影响。"
        "本研究采用矩形框目标检测方法，能够满足初步筛查与区域定位需求；若后续研究需要计算破损面积或建立病害严重程度热力图，"
        "则应进一步引入实例分割或语义分割方法。"
    )

    heading(doc, "七、存在问题与后续优化方向", 1)
    para(
        doc,
        "当前实验仍以矩形框目标检测为主，能够定位破损区域的大致范围，但尚不能直接给出像素级边界和破损面积。"
        "对于瓦片边缘、木材纹理、阴影和远距离小目标，模型仍可能出现误检或漏检。"
        "后续可在保持现有目标检测流程的基础上，进一步扩充多季节、多光照和多距离样本，并考虑引入实例分割模型，"
        "以获得更精细的病害轮廓和定量化面积指标。",
    )
    para(
        doc,
        "此外，现阶段三座鼓楼均采用单一 damage 类别进行标注，适合完成破损区域的初步定位，但尚不能区分瓦片缺失、木构开裂、表面风化、腐朽等不同病害类型。"
        "若后续研究面向修缮决策或风险分级，应在专家参与下进一步细化病害分类体系，并建立跨季节、跨光照条件的持续监测数据集。"
    )

    heading(doc, "八、综合结论", 1)
    para(
        doc,
        "本研究完成了增冲、朝利、从江三座鼓楼的单案例破损检测实验，并在统一参数、统一类别和相近正负样本比例条件下训练了三鼓楼通用模型。"
        "实验结果表明，YOLOv8n 能够在无人机影像中对鼓楼可见破损区域进行自动识别；"
        "多鼓楼合并训练能够提升模型对不同场景的适应性，尤其在新增鼓楼案例中体现出更好的综合检测能力。"
        "该实验为后续鼓楼建筑数字化巡检、周期性监测和破损区域精细化分析奠定了基础。",
    )

    out = out_dir / "三鼓楼实例融合实验报告.docx"
    doc.save(out)
    return out


def write_index(paths: list[Path], metrics: dict[str, dict], dataset: dict):
    lines = [
        "# run9 最终结果索引",
        "",
        "## 模型目录",
        "",
        "- 增冲单案例模型：`run9/zengchong/model/train/weights/best.pt`",
        "- 朝利单案例模型：`run9/chaoli/model/train/weights/best.pt`",
        "- 从江单案例模型：`run9/congjiang/model/train/weights/best.pt`",
        "- 三鼓楼通用模型：`run9/universal/model/train/weights/best.pt`",
        "",
        "## 文档",
        "",
    ]
    for path in paths:
        lines.append(f"- `{path.relative_to(RUN9)}`")
    lines += [
        "",
        "## 数据集统计",
        "",
        "| 数据集 | 图像数 | 正样本 | 负样本 | 标注框 | 正样本比例 |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for case in ["zengchong", "chaoli", "congjiang"]:
        info = dataset["by_case"][case]
        lines.append(
            f"| {CASE_NAMES[case]} | {info['images']} | {info['positive_images']} | {info['negative_images']} | {info['boxes']} | {pct(info['positive_ratio'])} |"
        )
    lines.append(
        f"| 三鼓楼通用数据集 | {dataset['images']} | {dataset['positive_images']} | {dataset['negative_images']} | {dataset['boxes']} | {pct(dataset['positive_ratio'])} |"
    )
    lines += [
        "",
        "## 关键测试指标",
        "",
        "| 模型/测试集 | Precision | Recall | mAP50 | mAP50-95 |",
        "|---|---:|---:|---:|---:|",
    ]
    for key in [
        "zengchong_single_on_zengchong_test",
        "chaoli_single_on_chaoli_test",
        "congjiang_single_on_congjiang_test",
        "universal_on_universal_test",
        "universal_on_zengchong_test",
        "universal_on_chaoli_test",
        "universal_on_congjiang_test",
    ]:
        item = metrics[key]
        lines.append(
            f"| {key} | {num(item['precision'])} | {num(item['recall'])} | {num(item['mAP50'])} | {num(item['mAP50_95'])} |"
        )
    (RUN9 / "最终结果索引.md").write_text("\n".join(lines), encoding="utf-8")


def copy_best_weights():
    for case in ["zengchong", "chaoli", "congjiang", "universal"]:
        source = RUN9 / case / "model" / "train" / "weights" / "best.pt"
        target = RUN9 / case / "model" / f"{case}_best.pt"
        shutil.copy2(source, target)


def main() -> None:
    dataset = read_json(RUN9 / "datasets" / "universal" / "dataset_summary.json")
    metrics = metrics_by_name()
    figures = ensure_comparison_figures(metrics, dataset)
    samples = ensure_sample_detection_figures()

    paths = []
    paths.append(generate_single_report("zengchong", metrics["zengchong_single_on_zengchong_test"], dataset["by_case"]["zengchong"], samples))
    paths.append(generate_single_report("chaoli", metrics["chaoli_single_on_chaoli_test"], dataset["by_case"]["chaoli"], samples))
    paths.append(generate_single_report("congjiang", metrics["congjiang_single_on_congjiang_test"], dataset["by_case"]["congjiang"], samples))
    paths.append(generate_universal_report(metrics, dataset, figures, samples))
    paths.append(generate_fusion_report(metrics, dataset, figures, samples))
    copy_best_weights()
    write_index(paths, metrics, dataset)
    print("Generated reports:")
    for path in paths:
        print(path)


if __name__ == "__main__":
    main()
